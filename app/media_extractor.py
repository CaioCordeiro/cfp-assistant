import io
import os
from typing import Optional

import requests
from requests.auth import HTTPBasicAuth
from pypdf import PdfReader
from pptx import Presentation


SUPPORTED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
}


def _twilio_auth() -> Optional[HTTPBasicAuth]:
    sid = os.environ.get("TWILIO_ACCOUNT_SID")
    token = os.environ.get("TWILIO_AUTH_TOKEN")
    if sid and token:
        return HTTPBasicAuth(sid, token)
    return None


def fetch_media(url: str) -> bytes:
    auth = _twilio_auth()
    resp = requests.get(url, auth=auth, timeout=10)
    resp.raise_for_status()
    return resp.content


def extract_text_from_pdf(data: bytes, max_chars: int = 6000) -> str:
    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt:
            chunks.append(txt.strip())
        if sum(len(c) for c in chunks) >= max_chars:
            break
    text = "\n\n".join(chunks)
    return text[:max_chars]


def extract_text_from_pptx(data: bytes, max_chars: int = 6000) -> str:
    prs = Presentation(io.BytesIO(data))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                t = (shape.text or "").strip()
                if t:
                    texts.append(t)
        if sum(len(t) for t in texts) >= max_chars:
            break
    text = "\n\n".join(texts)
    return text[:max_chars]


def extract_media_text(url: str, content_type: str) -> Optional[str]:
    kind = SUPPORTED_CONTENT_TYPES.get(content_type)
    if not kind:
        return None
    data = fetch_media(url)
    if kind == "pdf":
        return extract_text_from_pdf(data)
    if kind == "pptx":
        return extract_text_from_pptx(data)
    return None
